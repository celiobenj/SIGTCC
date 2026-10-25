"""Importação de alunos por arquivo CSV ou XLSX (spec 004; diagrama "Matrícula por CSV").

Leitura segura do arquivo (tamanho, tipo pelos bytes, zip bomb, XML seguro, limite de linhas),
validação de cada linha e prévia sem gravar nada (research B9 a B11).

[US-08][RBN-CAD-05.2][RBN-CAD-05.5][RBN-CAD-05.7][FR-016][FR-017][Q2][Q4][Q11][Q13]
"""

import csv
import io
import zipfile
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from sigtcc.apps.cadastro import mensagens, services_turmas
from sigtcc.apps.cadastro.choices import StatusConta
from sigtcc.apps.cadastro.forms_turmas import normalizar_texto, validar_dados_do_aluno
from sigtcc.apps.cadastro.models import Alunos, AlunosTurmas

TAMANHO_MAXIMO = 2 * 1024 * 1024
LIMITE_DESCOMPACTADO = 20 * 1024 * 1024
RAZAO_MAXIMA = 100
LIMITE_LINHAS = 200
CABECALHOS = {"matrícula": "matricula", "nome completo": "nome", "curso": "curso"}
EXEMPLO = ("2415080008", "Nome Completo do Aluno", "EST08MCPD")


class ErroModelo(Exception):
    """Arquivo fora do modelo: "Arquivo fora do modelo. Baixe o modelo de importação".

    [US-08][RBN-CAD-05.2][FR-016]
    """


@dataclass
class Linha:
    """Uma linha de dados do arquivo, com o número da linha no arquivo (cabeçalho = 1).

    [US-08][RBN-CAD-05.2][FR-017]
    """

    numero: int
    matricula: str
    nome: str
    curso: str


def _texto_da_celula(valor):
    """Célula do XLSX ou do CSV como texto; número inteiro sem ".0". [US-08][B9]"""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def _mapear_cabecalho(cabecalho):
    """Posição de cada um dos três cabeçalhos, em qualquer ordem, sem diferenciar maiúsculas.

    [RBN-CAD-05.2][FR-016]
    """
    posicoes = {}
    for indice, valor in enumerate(cabecalho):
        chave = CABECALHOS.get(normalizar_texto(_texto_da_celula(valor)).casefold())
        if chave and chave not in posicoes:
            posicoes[chave] = indice
    if len(posicoes) != len(CABECALHOS):
        raise ErroModelo
    return posicoes


def _linhas_de(registros):
    """Aplica o cabeçalho às linhas, ignora as vazias e para acima de 200.

    [US-08][RBN-CAD-05.2][Q13][B-3]
    """
    iterador = iter(registros)
    try:
        posicoes = _mapear_cabecalho(next(iterador))
    except StopIteration as erro:
        raise ErroModelo from erro
    linhas = []
    for numero, registro in enumerate(iterador, start=2):
        valores = {
            chave: normalizar_texto(_texto_da_celula(registro[i] if i < len(registro) else None))
            for chave, i in posicoes.items()
        }
        if not any(valores.values()):
            continue
        if len(linhas) == LIMITE_LINHAS:
            raise ErroModelo
        linhas.append(Linha(numero, valores["matricula"], valores["nome"], valores["curso"]))
    return linhas


def _ler_csv(conteudo):
    """CSV em UTF-8 (com ou sem BOM) ou Windows-1252, com ";" ou ",". [US-08][Q11][m20]"""
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252")
    if "\x00" in texto or not texto.strip():
        raise ErroModelo
    primeira = texto.splitlines()[0]
    try:
        separador = csv.Sniffer().sniff(primeira, delimiters=";,").delimiter
    except csv.Error:
        separador = ";" if primeira.count(";") >= primeira.count(",") else ","
    return _linhas_de(csv.reader(io.StringIO(texto), delimiter=separador))


def _conferir_zip(conteudo):
    """Recusa XLSX que não é ZIP ou que se expande demais (zip bomb). [US-08][B-3][N7]"""
    if not conteudo.startswith(b"PK\x03\x04"):
        raise ErroModelo
    with zipfile.ZipFile(io.BytesIO(conteudo)) as zipado:
        partes = zipado.infolist()
        if sum(parte.file_size for parte in partes) > LIMITE_DESCOMPACTADO:
            raise ErroModelo
        for parte in partes:
            if parte.compress_size and parte.file_size / parte.compress_size > RAZAO_MAXIMA:
                raise ErroModelo


def _ler_xlsx(conteudo):
    """XLSX: só a primeira planilha, em modo só leitura, com XML seguro (defusedxml).

    [US-08][Q11][B-3]
    """
    _conferir_zip(conteudo)
    livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        return _linhas_de(livro.worksheets[0].iter_rows(values_only=True))
    finally:
        livro.close()


def ler_arquivo(arquivo):
    """Lê o arquivo enviado e devolve as linhas, ou levanta ErroModelo.

    Até 2 MB; extensão .csv ou .xlsx; qualquer erro de leitura é "fora do modelo", nunca 500.
    [US-08][RBN-CAD-05.2][FR-016][Q4][Q11][Q13][B-3]
    """
    if arquivo is None or arquivo.size > TAMANHO_MAXIMO:
        raise ErroModelo
    nome = (arquivo.name or "").lower()
    try:
        conteudo = arquivo.read()
        if nome.endswith(".csv"):
            return _ler_csv(conteudo)
        if nome.endswith(".xlsx"):
            return _ler_xlsx(conteudo)
    except ErroModelo:
        raise
    except (
        csv.Error,
        UnicodeDecodeError,
        zipfile.BadZipFile,
        InvalidFileException,
        KeyError,
        ValueError,
        IndexError,
        OSError,
        TypeError,
    ) as erro:
        raise ErroModelo from erro
    raise ErroModelo


# =============================================================================================
# Prévia (diagrama: validarLinha, buscarCurso, marcarAlertaCurso, buscarPorMatricula, existe)
# =============================================================================================

NOVO = "Novo"
CADASTRADO = "Já cadastrado"
REJEITADO = "Rejeitado"


@dataclass
class ItemDaPrevia:
    """Uma linha classificada na prévia, com o motivo da recusa e os avisos.

    [US-08][RBN-CAD-05.2][FR-017]
    """

    numero: int
    matricula: str
    nome: str
    curso: str
    situacao: str
    motivo: str = ""
    avisos: list = field(default_factory=list)


@dataclass
class Previa:
    """Linhas classificadas e as contagens mostradas antes de confirmar. [FR-017]"""

    itens: list

    @property
    def validos(self):
        """Linhas Novo ou Já cadastrado. [FR-017]"""
        return [item for item in self.itens if item.situacao != REJEITADO]

    @property
    def rejeitados(self):
        """Linhas rejeitadas, com o motivo. [FR-017]"""
        return [item for item in self.itens if item.situacao == REJEITADO]

    @property
    def nao_reconhecidos(self):
        """Linhas válidas com código de curso não reconhecido. [RBN-CAD-05.7][FR-018]"""
        return sum(1 for item in self.validos if mensagens.CURSO_NAO_RECONHECIDO in item.avisos)


def _classificar(turma, linha, vistos):
    """Classifica uma linha sem gravar nada, como o diagrama "Matrícula por CSV".

    [US-08][RBN-CAD-05.1][RBN-CAD-05.5][RBN-CAD-05.7][FR-017][Q7][Q16]
    """
    item = ItemDaPrevia(linha.numero, linha.matricula, linha.nome, linha.curso, REJEITADO)
    dados, motivos = validar_dados_do_aluno(linha.matricula, linha.nome, linha.curso)
    if motivos:
        item.motivo = motivos[0]
        return item
    if dados["matricula"] in vistos:
        item.motivo = mensagens.ALUNO_JA_NA_TURMA
        return item
    vistos.add(dados["matricula"])
    item.matricula, item.nome, item.curso = dados["matricula"], dados["nome"], dados["curso"]
    aluno = Alunos.objects.select_related("alu").filter(alu_matricula=item.matricula).first()
    codigo = item.curso
    if aluno is not None:
        if AlunosTurmas.objects.filter(mat_tur=turma, mat_alu=aluno).exists():
            item.motivo = mensagens.ALUNO_JA_NA_TURMA
            return item
        recusa, _ = services_turmas._recusa_por_tccs(aluno, turma)
        if recusa is not None:
            item.motivo = recusa.mensagem
            return item
        if aluno.alu.usu_status in (StatusConta.INATIVA, StatusConta.EXPIRADA):
            item.avisos.append(mensagens.MT29_CONTA_NAO_ATIVA.format(status=aluno.alu.usu_status))
        codigo = aluno.alu_codigo_curso
        item.situacao = CADASTRADO
    else:
        item.situacao = NOVO
    curso = services_turmas.curso_do_codigo(codigo)
    if curso is None:
        item.avisos.append(mensagens.CURSO_NAO_RECONHECIDO)
    elif curso.pk != turma.tur_cur_id:
        item.avisos.append(
            mensagens.MT13_CURSO_DIVERGENTE.format(
                curso_aluno=curso.cur_nome, curso_turma=turma.tur_cur.cur_nome
            )
        )
    return item


def montar_previa(turma, linhas):
    """Prévia com os alunos novos, os já cadastrados, as linhas com erro e os avisos (FR-017).

    [US-08][RBN-CAD-05.2][FR-017]
    """
    vistos = set()
    return Previa([_classificar(turma, linha, vistos) for linha in linhas])


def gerar_modelo_csv():
    """Modelo de importação: CSV em UTF-8 com BOM e ";", cabeçalho e uma linha de exemplo.

    [US-08][RBN-CAD-05.2][FR-016][Q11]
    """
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(("Matrícula", "Nome Completo", "Curso"))
    escritor.writerow(EXEMPLO)
    return saida.getvalue().encode("utf-8-sig")
